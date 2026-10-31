# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 The OpenConex contributors
import datetime as dt

import pytest
from docx import Document

from openconex.connections.base_plate.check import PressureModel
from openconex.io.project import CustomCombinationInput, HSSSectionInput
from openconex.report.builder import build_base_plate_report, factor_expression
from openconex.report.docx import latex_to_text, write_docx
from openconex.report.html import latex_to_mathml, render_html
from openconex.report.pdf import pdf_available, write_pdf

TODAY = dt.date(2026, 10, 30)


def test_report_content(project):
    r = build_base_plate_report(project, TODAY)
    assert r.ok
    assert len(r.summary) == len(r.checks) == 7
    assert {b.check_id for b in r.checks} >= {"base_plate.bearing", "anchor.steel_tension"}
    assert all(b.lines for b in r.checks)
    worst = max(r.checks, key=lambda b: b.ratio)
    assert worst.ratio == pytest.approx(0.976, abs=0.001)


def test_report_figures(project):
    r = build_base_plate_report(project, TODAY)
    assert [figure.key for figure in r.figures] == ["plan", "elevation"]
    assert all(figure.svg.startswith("<?xml") for figure in r.figures)
    html = render_html(r)
    assert html.count("<figure>") == 2
    assert html.count("<svg") >= 2


@pytest.mark.parametrize(("lang", "title", "status"), [("es", "Memoria de cálculo", "Cumple"),
                                                       ("en", "Design calculation report", "OK")])
def test_html_language(project, lang, title, status):
    project.info.language = lang
    html = render_html(build_base_plate_report(project, TODAY))
    assert f"<title>{title}" in html
    assert f">{status}<" in html
    assert "<math" in html  # equations rendered as MathML, no internet needed
    assert "2026-10-30" in html


def test_units_follow_project(project):
    project.info.units = "SI_kN_m"
    html = render_html(build_base_plate_report(project, TODAY))
    assert "kN" in html and "kip" not in html


def test_report_flags_failures(project):
    project.base_plate.tp = "1 in"
    r = build_base_plate_report(project, TODAY)
    assert not r.ok
    assert any(row.status == "ng" for row in r.summary)


def test_custom_combinations_in_report(project):
    project.combinations.include_code = False
    project.combinations.custom = [
        CustomCombinationInput(name="C-1", factors={"D": 1.2, "L": 1.6}),
    ]
    r = build_base_plate_report(project, TODAY)
    assert [row.cells[0] for row in r.combinations] == ["C-1"]
    assert r.combinations[0].cells[3] == "Personalizada"
    assert {block.combo for block in r.checks} == {"C-1"}


def test_triangular_pressure_assumption_in_report(project):
    project.base_plate.pressure_model = PressureModel.TRIANGULAR
    r = build_base_plate_report(project, TODAY)
    assert any("triangular" in assumption.lower() for assumption in r.assumptions)


def test_factor_expression():
    assert factor_expression({"D": 1.2, "L": 1.6}) == "1.2·D + 1.6·L"
    assert factor_expression({"D": 0.6, "W": -0.6}) == "0.6·D - 0.6·W"
    assert factor_expression({"W": -1.0}) == "-1·W"


@pytest.mark.parametrize(("shape", "grade", "text"), [
    ("square", "ASTM A500 Gr. C (rectangular)", "0.95"),
    ("round", "ASTM A500 Gr. C (round)", "0.8"),
])
def test_hss_column_report(project, shape, grade, text):
    project.base_plate.column = HSSSectionInput(
        type="HSS", shape=shape, d="12 in", t="0.5 in", grade=grade)
    r = build_base_plate_report(project, TODAY)
    assert r.inputs[0].cells[1].startswith("HSS ")
    assert any("HSS" in assumption and text in assumption for assumption in r.assumptions)


def test_latex_helpers():
    assert latex_to_text(r"\phi R_n = \frac{a}{b}") == "φ R_n = (a)/(b)"
    assert latex_to_text(r"\sqrt{\frac{4 T_u b'}{\phi p F_u}}") == "√((4 T_u b')/(φ p F_u))"
    assert latex_to_text(r"M_{pl} \le 1.7 f'_c") == "M_pl ≤ 1.7 f'_c"
    assert str(latex_to_mathml("")) == ""
    assert "<math" in str(latex_to_mathml(r"\frac{1}{2}"))


def test_docx(tmp_path, project):
    path = write_docx(build_base_plate_report(project, TODAY), tmp_path / "r.docx")
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("5. Verificaciones") for h in headings)
    assert len(doc.tables) >= 5 + 7


@pytest.mark.skipif(not pdf_available(), reason="Qt WebEngine not installed (openconex[pdf])")
def test_pdf(tmp_path, project):
    path = write_pdf(build_base_plate_report(project, TODAY), tmp_path / "r.pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and data.count(b"/Type /Page") >= 3
