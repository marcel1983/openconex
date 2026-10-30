# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 The OpenConex contributors
"""Editable Word report (python-docx). Equations are written as readable text."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from openconex.i18n import tr
from openconex.report.builder import Report

__all__ = ["latex_to_text", "write_docx"]

_SYMBOLS = {r"\phi": "φ", r"\Omega": "Ω", r"\alpha": "α", r"\beta": "β", r"\delta": "δ",
            r"\rho": "ρ", r"\lambda": "λ", r"\psi": "ψ", r"\mu": "μ", r"\pi": "π",
            r"\theta": "θ", r"\le": "≤", r"\ge": "≥", r"\cdot": "·", r"\times": "×",
            r"\max": "max", r"\min": "min", r"\sin": "sin", r"\tfrac": r"\frac"}


def _brace(text: str, start: int) -> tuple[str, int]:
    """Content of the {...} group starting at ``start`` and the index after it."""
    depth = 0
    for i in range(start, len(text)):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:i], i + 1
    return text[start + 1:], len(text)


def latex_to_text(latex: str) -> str:
    """Plain text rendering of the small LaTeX subset used in the traces."""
    s = latex
    for cmd in (r"\left", r"\right", r"\,", r"\ ", r"\!"):
        s = s.replace(cmd, " " if cmd in (r"\,", r"\ ") else "")
    for key in sorted(_SYMBOLS, key=len, reverse=True):
        s = s.replace(key, _SYMBOLS[key])
    out, i = "", 0
    while i < len(s):
        if s.startswith(r"\frac", i):
            num, j = _brace(s, i + 5)
            den, k = _brace(s, j)
            out += f"({latex_to_text(num)})/({latex_to_text(den)})"
            i = k
        elif s.startswith(r"\sqrt", i):
            arg, j = _brace(s, i + 5)
            out += f"√({latex_to_text(arg)})"
            i = j
        elif s.startswith(r"\text", i):
            arg, j = _brace(s, i + 5)
            out += arg
            i = j
        elif s[i] in "_^" and i + 1 < len(s) and s[i + 1] == "{":
            arg, j = _brace(s, i + 1)
            out += ("_" if s[i] == "_" else "^") + latex_to_text(arg)
            i = j
        elif s[i] in "{}":
            i += 1
        else:
            out += s[i]
            i += 1
    return re.sub(r"\s+", " ", out).strip()


def _table(doc, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = str(text)


def write_docx(report: Report, path: str | Path) -> Path:
    def t(key: str, **kw) -> str:
        return tr(key, report.language, **kw)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(report.title, level=0)
    doc.add_paragraph(report.subtitle)
    _table(doc, [t("report.item"), t("report.value")], [list(r) for r in report.info])
    verdict = doc.add_paragraph()
    run = verdict.add_run(t("report.result_ok" if report.ok else "report.result_ng"))
    run.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x7F, 0x37) if report.ok else RGBColor(0xC6, 0x28, 0x28)

    doc.add_heading(f"1. {t('report.codes')}", level=1)
    for code in report.codes:
        doc.add_paragraph(code, style="List Bullet")
    doc.add_heading(t("report.assumptions"), level=2)
    for item in report.assumptions:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading(f"2. {t('report.inputs')}", level=1)
    _table(doc, [t("report.item"), t("report.value")], [r.cells for r in report.inputs])
    doc.add_heading(t("report.loads"), level=2)
    _table(doc, ["ID", t("report.item"), "P", "V", "M"], [r.cells for r in report.cases])

    doc.add_heading(f"3. {t('report.combinations_title')}", level=1)
    _table(doc, ["ID", t("report.method"), t("report.factors"), t("report.source")],
           [r.cells for r in report.combinations])

    doc.add_heading(f"4. {t('report.summary')}", level=1)
    _table(doc, [t("report.check"), t("report.combination"), t("report.demand"),
                 t("report.capacity"), t("report.ratio"), t("report.status_col")],
           [[*r.cells, t(f"status.{r.status}")] for r in report.summary])

    doc.add_heading(f"5. {t('report.checks')}", level=1)
    for n, block in enumerate(report.checks, 1):
        doc.add_heading(f"5.{n} {block.title} — {t('status.' + block.status)}", level=2)
        doc.add_paragraph(f"{t('report.clause')}: {block.clause} · "
                          f"{t('report.combination')}: {block.combo}").italic = True
        rows = []
        for line in block.lines:
            eq = latex_to_text(line.symbol)
            if line.expression:
                eq += " = " + latex_to_text(line.expression)
            rows.append([eq, line.value, line.clause])
        rows.append([f"{t('report.demand')} / {t('report.capacity')}",
                     f"{block.demand} / {block.capacity} = {block.ratio:.3f}", ""])
        _table(doc, [t("report.item"), t("report.value"), t("report.clause")], rows)

    foot = doc.add_paragraph(f"{report.footer} {report.disclaimer}")
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in foot.runs:
        r.font.size = Pt(8)
    path = Path(path)
    doc.save(path)
    return path
